"""One-page board report (Word) and full data export (Excel)."""

from __future__ import annotations

import io
import re
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from montecarlo import export_excel, schedule_sensitivity
from montecarlo.integrated import as_result

from .engine import Answer, drivers, headline

NAVY = RGBColor(0x10, 0x42, 0x81)


def _shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    pr.append(sh)


def _table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); r.font.size = Pt(9)
        _shade(c, "104281")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def _chart(a: Answer) -> io.BytesIO:
    cur = a.project.currency
    has_s = a.res.schedule is not None
    fig, axes = plt.subplots(1, 2 if has_s else 1, figsize=(7.2 if has_s else 3.8, 2.3))
    axes = np.atleast_1d(axes)
    for ax in axes:
        for s in ("top", "right"):
            ax.spines[s].set_visible(False)
        ax.tick_params(labelsize=7, colors="#52514e"); ax.grid(axis="y", color="#e6e5e0", lw=.6); ax.set_axisbelow(True)
    xs = np.sort(a.res.total_cost)
    axes[0].plot(xs, np.linspace(0, 1, len(xs)), color="#2a78d6", lw=2)
    axes[0].axvline(a.base_cost, color="#52514e", ls=":", lw=1); axes[0].axhline(a.project.options.confidence, color="#9a9993", ls="--", lw=.8)
    axes[0].set_title("Chance cost is at or under", fontsize=8, loc="left")
    axes[0].xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{cur}{v/1000:,.0f}k"))
    axes[0].yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
    if has_s:
        d = np.sort(a.res.schedule.project_days)
        axes[1].plot(d, np.linspace(0, 1, len(d)), color="#1baf7a", lw=2)
        axes[1].axvline(a.plan_days, color="#52514e", ls=":", lw=1); axes[1].axhline(a.project.options.confidence, color="#9a9993", ls="--", lw=.8)
        axes[1].set_title("Chance of finishing within (working days)", fontsize=8, loc="left")
        axes[1].yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
    buf = io.BytesIO(); fig.tight_layout(); fig.savefig(buf, format="png", dpi=220); plt.close(fig); buf.seek(0)
    return buf


def board_report(a: Answer, author: str = "") -> bytes:
    p, cur = a.project, a.project.currency
    conf = int(p.options.confidence * 100)
    dv = drivers(a)
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.8); s.top_margin = s.bottom_margin = Cm(1.5)
    h = doc.add_heading(f"{p.name}: cost and schedule confidence", 1); h.runs[0].font.color.rgb = NAVY
    meta = doc.add_paragraph(f"RiskSim · {p.options.n_sims:,} simulations · start {p.start_date}" + (f" · prepared by {author}" if author else ""))
    meta.runs[0].italic = True; meta.runs[0].font.size = Pt(8.5)

    para = doc.add_paragraph()
    for i, part in enumerate(re.split(r"\*\*", headline(a))):
        run = para.add_run(part); run.bold = i % 2 == 1

    rows = [["Budget", f"{cur}{a.base_cost:,.0f}", f"{cur}{a.p50_cost:,.0f}", f"{cur}{a.budget:,.0f}",
             f"{cur}{a.contingency:,.0f}"]]
    if a.date is not None:
        rows.append(["Finish date", f"{a.plan_date:%d %b %Y}", f"{a.p50_date:%d %b %Y}", f"{a.date:%d %b %Y}",
                     f"{a.date_days - a.plan_days:.0f} working days"])
    _table(doc, ["", "Plan (no uncertainty)", "P50 (coin flip)", f"P{conf} (recommended)", "Allowance needed"], rows,
           [3, 3.6, 3.2, 3.6, 3.6])
    doc.add_paragraph()
    doc.add_picture(_chart(a), width=Cm(17.4 if a.res.schedule is not None else 9))

    sp = dv["split"]
    doc.add_paragraph().add_run("Where the contingency comes from").bold = True
    doc.add_paragraph(f"Estimating uncertainty {cur}{sp['estimate_uncertainty']:,.0f} · named risks "
                      f"{cur}{sp['risk_events']:,.0f}. The named-risk share is what risk responses can reduce.")
    if len(dv["risks"]):
        doc.add_paragraph().add_run("Risks to act on first").bold = True
        rk = dv["risks"]
        cc = [c for c in rk.columns if c.endswith("cost saved if removed")][0]
        dc = [c for c in rk.columns if c.endswith("days saved if removed")]
        resp = dict(zip(p.risks["ID"].astype(str), p.risks["Response"].fillna("")))
        _table(doc, ["Risk", "Owner", "Chance", f"P{conf} cost if removed", "Days if removed", "Planned response"],
               [[f"{r['Risk ID']} {r['Risk']}", r["Owner"], f"{r['Probability']:.0%}", f"−{cur}{r[cc]:,.0f}",
                 (f"−{r[dc[0]]:.0f}" if r[dc[0]] >= 0.5 else "0") if dc else "–", resp.get(r["Risk ID"], "")] for _, r in rk.iterrows()],
               [4.6, 2.4, 1.4, 2.4, 1.8, 4.8])
    if len(dv["critical"]):
        doc.add_paragraph()
        doc.add_paragraph().add_run("Tasks that control the end date").bold = True
        crit = dv["critical"]
        doc.add_paragraph("; ".join(f"{r['Task']} (critical in {r['Criticality index']:.0%} of runs)"
                                    for _, r in crit.head(3).iterrows()) + ".")
    note = doc.add_paragraph(f"Method: Monte Carlo simulation of three-point estimates ({p.options.distribution}), "
                             "named risks as probabilistic events, critical-path scheduling on every run.")
    note.runs[0].italic = True; note.runs[0].font.size = Pt(8)
    out = io.BytesIO(); doc.save(out)
    return out.getvalue()


def excel_report(a: Answer) -> bytes:
    dv = drivers(a)
    extras = {"Contingency allocation": dv["allocation"]}
    if a.res.schedule is not None:
        extras["Schedule"] = schedule_sensitivity(a.res.schedule)
    if len(dv["risks"]):
        extras["Risk ranking"] = dv["risks"]
    tmp = Path(tempfile.mkdtemp()) / "risksim.xlsx"
    export_excel({a.res.distribution: as_result(a.res)}, a.res.estimates, tmp, extras)
    return tmp.read_bytes()
